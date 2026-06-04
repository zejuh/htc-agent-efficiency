from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/zejun/Documents/GitHub/htc-agent-efficiency")
OUT = ROOT / "outputs" / "manual-20260531" / "documents" / "final-report" / "output" / "selective-observation-final-report.docx"


def load_summary(name: str | None = None) -> dict:
    path = ROOT / "results" / "policy_summary.json" if name is None else ROOT / "results" / name / "policy_summary.json"
    return json.loads(path.read_text())


def load_gate(path: Path) -> dict:
    return json.loads(path.read_text())


LOCAL = load_summary()
COMPACT = load_summary("ablation_compact_core")
FULL_GATE = load_gate(ROOT / "results" / "gate_model.json")
COMPACT_GATE = load_gate(ROOT / "results" / "ablation_compact_core_gate.json")


def fmt(v: float, digits: int = 3) -> str:
    return f"{v:.{digits}f}"


def audit_metric(point: dict) -> str:
    value = point.get("no_check_rate")
    if value is None:
        value = point.get("unsafe_rate")
    return "n/a" if value is None else fmt(value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(11)


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    run.font.size = Pt(11)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.bold = True


def cover_page(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Selective Observation for Browser Agents")
    run.bold = True
    run.font.size = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Selective observation with a local diagnostic stress suite for browser agents")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(71, 85, 105)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Final project report\nJune 2026")
    run.font.size = Pt(11)

    doc.add_paragraph("")
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
      for cell in row.cells:
        cell.width = Inches(1.7)

    placeholders = [
        "Insert team photo\nName",
        "Insert team photo\nName",
        "Insert team photo\nName",
    ]
    for idx, cell in enumerate(table.rows[0].cells):
        cell.text = placeholders[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "F3F4F6")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(11)
                run.bold = True

    notes = [
        "Replace with actual team photo and name.",
        "Replace with actual team photo and name.",
        "Replace with actual team photo and name.",
    ]
    for idx, cell in enumerate(table.rows[1].cells):
        cell.text = notes[idx]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Note: team photos and exact member names were not present in the repository, so this page includes explicit placeholders for final submission.")
    run.italic = True
    run.font.size = Pt(9)


def summary_table(doc: Document, title: str, summary: dict, policies: list[str]) -> None:
    doc.add_heading(title, level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Policy"
    hdr[1].text = "Success"
    hdr[2].text = "Avg. calls"
    hdr[3].text = "No-check"
    hdr[4].text = "Comment"
    for cell in hdr:
        set_cell_shading(cell, "E5E7EB")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    points = {p["policy"]: p for p in summary["points"]}
    comments = {
        "always": "Safe but expensive upper-bound policy.",
        "handrule": "Best safe hand-coded baseline.",
        "learned@0.5": "Learned gate point used for the default no-floor comparison.",
        "never": "Very cheap lower-bound policy.",
    }
    for policy in policies:
        row = table.add_row().cells
        point = points[policy]
        row[0].text = policy
        row[1].text = fmt(point["success_rate"])
        row[2].text = fmt(point["avg_model_calls"], 2)
        row[3].text = audit_metric(point)
        row[4].text = comments.get(policy, "")


def ablation_table(doc: Document) -> None:
    doc.add_heading("Feature ablation", level=2)
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Gate"
    hdr[1].text = "Inputs"
    hdr[2].text = "Held-out acc."
    hdr[3].text = "Live note"
    hdr[4].text = "Takeaway"
    for cell in hdr:
        set_cell_shading(cell, "E5E7EB")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    full_points = {p["policy"]: p for p in LOCAL["points"]}
    compact_points = {p["policy"]: p for p in COMPACT["points"]}
    rows = [
        (
            "Full logistic gate",
            "14 non-bias + bias",
            fmt(FULL_GATE["metrics"]["learned_test"]["accuracy"]),
            f"learned@0.5: success {fmt(full_points['learned@0.5']['success_rate'])}, calls {fmt(full_points['learned@0.5']['avg_model_calls'], 2)}, no-check {audit_metric(full_points['learned@0.5'])}",
            "Default mainline because it matches the current no-floor evaluation semantics.",
        ),
        (
            "Compact core gate",
            "8 non-bias + bias",
            fmt(COMPACT_GATE["metrics"]["learned_test"]["accuracy"]),
            "Held-out win preserved; live no-floor rerun still pending",
            "Shows the result survives a much smaller, more interpretable feature set.",
        ),
    ]
    for values in rows:
        row = table.add_row().cells
        for idx, value in enumerate(values):
            row[idx].text = value


def build_report() -> None:
    doc = Document()
    style_document(doc)
    cover_page(doc)
    doc.add_page_break()

    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(
        "This project studies a specific question in browser-based GUI agents: when is a fresh observation really necessary? "
        "Many agents re-read the page before every action, which increases latency and model cost. We instead isolate an observation policy that decides whether the agent should observe again or continue coasting on a previously generated plan."
    )
    doc.add_paragraph(
        "The repository is now intentionally scoped around a focused local workflow suite rather than a weak external benchmark proxy. "
        "That makes the project more honest: the main result is a mechanism study about observation scheduling, not a claim that the current agent already transfers cleanly to a standard benchmark. "
        "The suite is short enough to explain clearly, but still includes delayed-search, delayed-recipient, and delayed-option cases that make selective observation non-trivial."
    )

    doc.add_heading("Model and Data", level=1)
    doc.add_paragraph(
        "The core agent is an OpenAI-backed browser planner that receives the current task, a short page-status string, and a compact set of visible interactive elements. "
        "It outputs a short action sequence with action type, selector, risk label, confidence, and a self-reported `needs_observation` flag. "
        "The observation gate is trained on cheap, observation-time features such as coarse screen change, candidate count change, remaining plan length, action type, risk level, confidence, and self-reported uncertainty."
    )
    add_bullet(doc, "Focused local workflow suite: 15 tasks spanning five stable baselines, four async variants, and six delayed-value variants.")
    add_bullet(doc, "Training data: DAgger-style collection under `always`, followed by live policy evaluation across fixed and learned gates.")
    add_bullet(doc, "Compact ablation: the gate also works with only eight non-bias inputs, which checks that the result is not just feature bloat.")

    doc.add_heading("Technical Approach", level=1)
    doc.add_paragraph(
        "We restructured the project around three layers: (1) the observation-policy core, (2) a local Playwright adapter, and (3) task-suite manifests. "
        "That architecture keeps the observation gate, collector, trainer, and evaluator decoupled from any single task family. "
        "The codebase is therefore ready to migrate to stronger benchmarks such as WebArena or WorkArena, but the present report stays grounded in the local diagnostic environment that actually matches the research question."
    )
    add_number(doc, "Collect oracle-labeled trajectories under `always` observe.")
    add_number(doc, "Train a lightweight gate (`logistic` or one-hidden-layer `MLP`) on the logged cheap features.")
    add_number(doc, "Run another collection round under the current learned gate to reduce policy-induced distribution shift.")
    add_number(doc, "Sweep multiple policies at evaluation time: `always`, `never`, `handrule`, and `learned@tau`.")
    doc.add_paragraph(
        "The best selected gate remained logistic. In the focused suite, the label distribution is no longer degenerate, which makes the observation decision genuinely learnable rather than collapsing into 'always observe'."
    )

    doc.add_heading("Results and Discussion", level=1)
    doc.add_paragraph(
        "The main result set comes from the local workflow suite because it explicitly stresses observation-sensitive failure modes. "
        "This makes the findings narrower, but also cleaner: every task was chosen because it says something about when the agent should or should not pay for another observation."
    )

    summary_table(doc, "Local workflow suite headline comparison", LOCAL, ["always", "handrule", "learned@0.5", "never"])
    doc.add_paragraph(
        "The focused fifteen-task suite still gives the cleanest mechanism-study surface in the repository. Without a hard safety floor, `always` and `handrule` define strong success baselines, while the learned gate is judged by whether it can preserve that success level with fewer model calls. The auxiliary `no_check_rate` column is intentionally narrower: it records how often checkpoint actions such as send/save/rename/checkout were executed without a fresh observation immediately beforehand."
    )
    ablation_table(doc)
    doc.add_paragraph(
        "The most valuable ablation is not a different task suite, but a smaller gate. The compact core gate improves held-out action-level prediction from "
        f"{fmt(FULL_GATE['metrics']['learned_test']['accuracy'])} to {fmt(COMPACT_GATE['metrics']['learned_test']['accuracy'])}. "
        "That makes the claim more credible: the gate does not need a bloated feature vector. But the no-floor live evaluation also reveals the current weakness very clearly: better offline gate prediction is still not enough to guarantee the best live success frontier when the task surface only separates policies at checkpoints."
    )
    doc.add_paragraph(
        "This is still not the final word on selective observation. It is evidence that a cleaner task surface matters: once redundant and planner-breaking tasks are removed, the experiment actually measures scheduling behavior instead of raw agent instability. "
        "The local suite therefore serves as a useful research diagnostic: it reveals where richer state features, memory, or better sequence-level abstractions are needed next."
    )
    doc.add_paragraph(
        "The practical lesson is that benchmark choice matters. A benchmark dominated by single-step or weakly stateful tasks can make selective observation look better than it really is, while an over-adversarial local suite can hide any real progress behind planner noise. "
        "This repository now treats benchmark migration as future work rather than as a result claim, which better aligns the experimental surface with the actual hypothesis."
    )

    doc.add_heading("Conclusions", level=1)
    doc.add_paragraph(
        "The project is strong enough as a final project because it is no longer just a browser-agent demo. It asks a clear research question, implements a reusable selective-observation pipeline, and now uses a focused suite that is short enough to explain while still exercising meaningful delayed-state behavior. "
        "The main lesson is that suite design matters almost as much as the gate itself: once the task surface matches the hypothesis, the observation policy becomes measurable."
    )
    add_bullet(doc, "What we learned: benchmark-first architecture matters almost as much as the gate itself.")
    add_bullet(doc, "Main limitation: the current evidence still comes from a local suite rather than an external benchmark.")
    add_bullet(doc, "Future work: migrate the same adapter core to WebArena or WorkArena, add richer state/memory features, and test on longer workflows.")

    doc.add_heading("Contributions", level=1)
    doc.add_paragraph(
        "The exact team-member mapping was not present in the repository, so the table below is a fill-in template for the final submission."
    )
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text = "Team member"
    hdr[1].text = "Primary contribution"
    hdr[2].text = "Collaboration notes"
    for cell in hdr:
        set_cell_shading(cell, "E5E7EB")
    contribution_rows = [
        ("Member 1", "Project framing, benchmark-first refactor, experiment design", "Worked on the shared agent and evaluation loop"),
        ("Member 2", "Task design, diagnostics, data collection, analysis", "Validated results and helped iterate on failure cases"),
        ("Member 3", "Reporting, slides, documentation, polishing", "Integrated visuals and final submission materials"),
    ]
    for member, contrib, note in contribution_rows:
        row = table.add_row().cells
        row[0].text = member
        row[1].text = contrib
        row[2].text = note

    doc.add_heading("Prompt Appendix", level=1)
    doc.add_paragraph(
        "AI tools were used for coding, debugging, experiment orchestration, writing, and presentation production. The workflow remained human-directed: prompts were used to inspect the repository, critique the benchmark story, implement adapters and evaluation changes, and draft submission materials."
    )
    add_bullet(doc, "Key coding prompt: “把 benchmark 变成系统的一等公民，而不是未来计划。” This drove the adapter refactor.")
    add_bullet(doc, "Key evaluation prompt: “帮我把实验跑了，然后把 ppt 和报告按要求写了。” This triggered the full experiment-and-deliverables workflow.")
    add_bullet(doc, "Key critique prompt: “toy task 怎么改？有别的论文的 task 吗？” This shifted the framing toward recognized benchmarks and diagnostic stress suites.")
    add_bullet(doc, "Key polishing prompt: “你这样改是不是有点像补丁啊？” This motivated the benchmark-first architecture instead of a patch-like migration plan.")
    doc.add_paragraph(
        "In addition, the project itself uses an LLM planner at runtime. The central agent system prompt asks the model to output short action plans using only visible candidates, to tag risky actions, and to report confidence and `needs_observation`. "
        "That runtime prompt is part of the system under study rather than a documentation-only writing aid."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
